"""入出金管理表 Excel への取引追記。

既存ヘッダ行を読み、設定で指定された列名に取引内容をマッピングして末尾に追記。
重複防止のため (日付, 摘要, 金額, 銀行) の組をキーに既存行と突合する。
"""
from __future__ import annotations

from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from pdf_parser import Transaction


def _row_key(date_val, desc, deposit, withdraw, bank) -> tuple:
    return (
        date_val.isoformat() if hasattr(date_val, "isoformat") else str(date_val),
        (desc or "").strip(),
        deposit or 0,
        withdraw or 0,
        (bank or "").strip(),
    )


def append_transactions(xlsx_path: Path, sheet_name: str, header_row: int,
                        column_map: dict[str, str],
                        transactions: list[Transaction]) -> int:
    """取引を Excel に追記。戻り値は追記件数(重複スキップを除く)。"""
    wb = load_workbook(xlsx_path)
    if sheet_name not in wb.sheetnames:
        raise ValueError(f"シート '{sheet_name}' が見つかりません: {wb.sheetnames}")
    ws = wb[sheet_name]

    headers = {ws.cell(row=header_row, column=c).value: c
               for c in range(1, ws.max_column + 1)
               if ws.cell(row=header_row, column=c).value}

    # ヘッダ → 列番号 の確認
    field_to_col: dict[str, int] = {}
    for header_name, field in column_map.items():
        if header_name not in headers:
            raise ValueError(
                f"Excelヘッダに '{header_name}' が見つかりません。"
                f"既存ヘッダ: {list(headers)}"
            )
        field_to_col[field] = headers[header_name]

    # 既存行のキー集合
    existing: set[tuple] = set()
    date_col = field_to_col.get("date")
    desc_col = field_to_col.get("description")
    dep_col = field_to_col.get("deposit")
    wd_col = field_to_col.get("withdraw")
    bank_col = field_to_col.get("bank")

    for r in range(header_row + 1, ws.max_row + 1):
        d = ws.cell(r, date_col).value if date_col else None
        if isinstance(d, datetime):
            d = d.date()
        desc = ws.cell(r, desc_col).value if desc_col else None
        dep = ws.cell(r, dep_col).value if dep_col else None
        wd = ws.cell(r, wd_col).value if wd_col else None
        bank = ws.cell(r, bank_col).value if bank_col else None
        if d:
            existing.add(_row_key(d, desc, dep, wd, bank))

    write_row = ws.max_row + 1
    appended = 0
    for tx in transactions:
        key = _row_key(tx.date, tx.description, tx.deposit, tx.withdraw, tx.bank)
        if key in existing:
            continue
        for field, col in field_to_col.items():
            value = {
                "date": tx.date,
                "description": tx.description,
                "deposit": tx.deposit,
                "withdraw": tx.withdraw,
                "balance": tx.balance,
                "bank": tx.bank,
                "note": tx.source_file,
            }.get(field)
            ws.cell(row=write_row, column=col, value=value)
        existing.add(key)
        write_row += 1
        appended += 1

    wb.save(xlsx_path)
    return appended
